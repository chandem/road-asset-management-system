"""Excel export helpers for RAMS reports."""

from io import BytesIO

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font


def _write_sheet(ws, headers: list[str], rows: list[dict]):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h) for h in headers])


def build_reports_workbook(maint: dict, condition: dict, defects: dict, costs: dict) -> StreamingResponse:
    wb = Workbook()

    ws_m = wb.active
    ws_m.title = "Maintenance"
    maint_headers = [
        "maintenance_id", "road_code", "road_name", "section_id", "activity_type",
        "priority", "planned_date", "completed_date", "status", "estimated_cost",
        "actual_cost", "contractor", "description",
    ]
    _write_sheet(ws_m, maint_headers, maint.get("items") or [])
    ws_m.append([])
    ws_m.append(["Summary", "activity_count", "completed_count", "estimated_cost", "actual_cost"])
    s = maint.get("summary") or {}
    ws_m.append(["", s.get("activity_count"), s.get("completed_count"), s.get("estimated_cost"), s.get("actual_cost")])

    ws_c = wb.create_sheet("Condition")
    _write_sheet(
        ws_c,
        ["road_code", "road_name", "section_count", "average_condition", "minimum_condition", "maximum_condition"],
        condition.get("items") or [],
    )

    ws_d = wb.create_sheet("Defects")
    _write_sheet(
        ws_d,
        [
            "defect_id", "road_code", "road_name", "defect_type", "severity",
            "chainage_km", "length_m", "width_m", "depth_mm", "detected_by",
        ],
        defects.get("items") or [],
    )
    ws_d.append([])
    ds = defects.get("summary") or {}
    ws_d.append(["Total defects", ds.get("defect_count")])

    ws_cost = wb.create_sheet("Costs")
    _write_sheet(
        ws_cost,
        ["road_code", "road_name", "activity_count", "estimated_cost", "actual_cost", "variance"],
        costs.get("items") or [],
    )
    ws_cost.append([])
    cs = costs.get("summary") or {}
    ws_cost.append(["budget", cs.get("budget")])
    ws_cost.append(["estimated_cost", cs.get("estimated_cost")])
    ws_cost.append(["actual_cost", cs.get("actual_cost")])
    ws_cost.append(["remaining_budget", cs.get("remaining_budget")])
    ws_cost.append(["budget_utilization_percent", cs.get("budget_utilization_percent")])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=rams-reports.xlsx"},
    )
